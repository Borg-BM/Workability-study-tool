"""Generates the first draft of a study's workability report as a Barge Master-style .docx.

Structure and standard texts follow BM's workability report templates (BM-QU-TE-5404 W2W
operation, 5405/5406 lifting operation) and the reports issued from them: executive summary
with an overview of results, introduction, input parameters, workability results, references,
Appendix A (results) and Appendix B (workability criterion). The page layout is the Memo
template's (BM-QU-TE-5204) that those templates build on: an uppercase title with the
Subject/Project/Ref./... block and the logo top right, a blue rule, numbered uppercase blue
headings, "Ref. - Subject" plus date as the running header, and a blue-ruled
"Barge Master  n/N" footer on every page.
"""

import base64
import io
import json
import re
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image

from . import db

# House style, logos and page setup from BM's Word templates (via the bm-document-generator skill,
# with the body styles checked against the Memo template).
_STYLE_DIR = Path(__file__).resolve().parent / "report_style"
_HOUSE_STYLE = json.loads((_STYLE_DIR / "bm_house_style.json").read_text(encoding="utf-8"))
_LOGO_FULL = json.loads((_STYLE_DIR / "bm_logo_full.json").read_text(encoding="utf-8"))
_LOGO_MARK = json.loads((_STYLE_DIR / "bm_logo_mark.json").read_text(encoding="utf-8"))
# Figure 1 of every BM workability report: the vessel axis/heading convention (DIN 81209-1).
_AXIS_CONVENTION = _STYLE_DIR / "axis_convention.png"

_PAGE = _HOUSE_STYLE["page_setup"]
_CONTENT_WIDTH = Pt(_PAGE["page_width_pt"] - _PAGE["margin_left_pt"] - _PAGE["margin_right_pt"])
_HEADER_FILL = _HOUSE_STYLE["table_style"]["header_shading"]["fill"]
_BODY_FILL = _HOUSE_STYLE["table_style"]["body_shading"][0]["fill"]

_BM_BLUE = "00B0F0"  # rules under the title block and above the footer
_FRAME_COLOR = "002060"  # thin frame around the figures
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_PLACEHOLDER_COLOR = RGBColor(0x88, 0x88, 0x88)

# Also in src/reportModel.ts, which makes the same default subject line for the setup form.
SYSTEM_LABELS = {"mcg30": "BM-MCG", "t700": "BM-T700", "heavy_feeder": "BM-HF"}
# How the reports name the system in running text: in full once, in the introduction...
_SYSTEM_FULL_NAMES = {
    "mcg30": "Barge Master (BM) Motion Compensated Gangway (MCG)",
    "t700": "Barge Master T700 platform (BM-T700)",
    "heavy_feeder": "Barge Master Heavy Feeder (BM-HF)",
}
# ...and short everywhere else.
_SYSTEM_SHORT_NAMES = {"mcg30": "MCG", "t700": "BM-T700", "heavy_feeder": "BM-HF"}

# Engineered lift capacity of the MCG30 lifting variants (see pybm_service.MACHINE_TYPES).
_LIFT_CAPACITY_KG = {"highload": 1000, "lowload": 500}

# The headings BM's reports summarise: bow waves and the +/-30 degree headings either side.
_SUMMARY_HEADINGS = ((180, "Head seas (180°)"), (210, "+30° heading (210°)"), (150, "-30° heading (150°)"))

_DNV_REFERENCE = "DNV-RP-C205 Environmental Conditions and Environmental Loads, Section 3.5.5.5, APRIL 2014"

_PROMPTS = {
    "executive_summary": "Conclude what the results mean for the operation, e.g. how suitable the vessel and location are.",
    "introduction_narrative": "Describe the assumptions behind this estimation (e.g. the arrangement and orientation the results are presented for).",
    "estimation_narrative": "Describe the figure below: which headings/periods are most/least restrictive, and what governs the limit.",
}


class _Prompt(str):
    """A [bracketed] writing prompt for the author, inside otherwise standard text."""


# ---- House style ----


def _apply_font(font, cfg: dict | None):
    if not cfg:
        return
    if cfg.get("name"):
        font.name = cfg["name"]
        # python-docx's base template sets theme fonts, which Word prefers over an explicit name.
        rfonts = font._element.get_or_add_rPr().rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
    if cfg.get("size_pt"):
        font.size = Pt(cfg["size_pt"])
    for attr in ("bold", "italic", "underline"):
        if cfg.get(attr) is not None:
            setattr(font, attr, cfg[attr])
    if cfg.get("caps") is not None:
        font.all_caps = cfg["caps"]
    if cfg.get("color"):
        font.color.rgb = RGBColor.from_string(cfg["color"])


def _apply_paragraph_format(pf, cfg: dict | None):
    if not cfg:
        return
    if cfg.get("alignment"):
        pf.alignment = getattr(WD_ALIGN_PARAGRAPH, cfg["alignment"])
    for key, attr in (
        ("space_before_pt", "space_before"),
        ("space_after_pt", "space_after"),
        ("left_indent_pt", "left_indent"),
        ("right_indent_pt", "right_indent"),
        ("first_line_indent_pt", "first_line_indent"),
    ):
        if cfg.get(key) is not None:
            setattr(pf, attr, Pt(cfg[key]))
    spacing = cfg.get("line_spacing")
    if spacing is not None:
        # The template stores exact spacing in EMU (203200 = 16pt); small values are line multiples.
        pf.line_spacing = Emu(int(spacing)) if spacing > 10 else spacing
        if cfg.get("line_spacing_rule"):
            pf.line_spacing_rule = getattr(WD_LINE_SPACING, cfg["line_spacing_rule"])
    for position in cfg.get("tab_stops_pt") or []:
        pf.tab_stops.add_tab_stop(Pt(position))
    if cfg.get("keep_with_next") is not None:
        pf.keep_with_next = cfg["keep_with_next"]


def _add_style(document, name: str, *, font: dict | None = None, paragraph_format: dict | None = None):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    _apply_font(style.font, font)
    _apply_paragraph_format(style.paragraph_format, paragraph_format)


def _apply_house_style(document):
    section = document.sections[0]
    section.page_width = Pt(_PAGE["page_width_pt"])
    section.page_height = Pt(_PAGE["page_height_pt"])
    section.top_margin = Pt(_PAGE["margin_top_pt"])
    section.bottom_margin = Pt(_PAGE["margin_bottom_pt"])
    section.left_margin = Pt(_PAGE["margin_left_pt"])
    section.right_margin = Pt(_PAGE["margin_right_pt"])

    # python-docx's base template defaults every run's East Asian font to Japanese MS Mincho.
    default_fonts = document.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts"))
    default_fonts.attrib.clear()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        default_fonts.set(qn(attr), _HOUSE_STYLE["default_font"]["name"])
    theme_lang = document.settings.element.find(qn("w:themeFontLang"))
    if theme_lang is not None:
        theme_lang.attrib.pop(qn("w:eastAsia"), None)

    paragraph_styles = {name: cfg for name, cfg in _HOUSE_STYLE["styles"].items() if cfg["type"] == "PARAGRAPH"}
    existing = {s.name for s in document.styles}
    for name in paragraph_styles:
        if name not in existing:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    for name, cfg in paragraph_styles.items():
        style = document.styles[name]
        if cfg.get("based_on") and cfg["based_on"] != name and cfg["based_on"] in paragraph_styles:
            style.base_style = document.styles[cfg["based_on"]]
        _apply_font(style.font, cfg.get("font"))
        _apply_paragraph_format(style.paragraph_format, cfg.get("paragraph_format"))

    # The few paragraph looks the reports use besides the template's own styles. As styles rather
    # than direct formatting, so text the app later rewrites (Check, Ref.) keeps the look.
    compact = {"alignment": "LEFT", "space_before_pt": 0.0, "space_after_pt": 0.0}
    _add_style(document, "BM Title Block", font={"name": "Arial", "size_pt": 8.0}, paragraph_format=compact)
    _add_style(document, "BM Running Header", font={"name": "Arial", "size_pt": 8.0}, paragraph_format=compact)
    _add_style(document, "BM Table Text", paragraph_format=compact)
    for name in ("Header", "Footer"):
        style = document.styles[name].paragraph_format
        _apply_paragraph_format(style, compact)
        style.tab_stops.clear_all()  # python-docx's are for Letter paper; the footer sets its own


def _add_logo(paragraph, logo: dict, *, width: int | None = None, height: int | None = None):
    run = paragraph.add_run()
    run.add_picture(
        io.BytesIO(base64.b64decode(logo["data_base64"])),
        width=Emu(width or logo["width_emu"]),
        height=Emu(height or logo["height_emu"]),
    )
    return run


def _float(run, *, from_left_margin: int, from_page_top: int):
    """Turns the run's inline picture into one floating at a fixed spot (python-docx only
    makes inline pictures), as the templates place the title-page logo."""
    inline = run._r.find(".//" + qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300", "simplePos": "0",
        "relativeHeight": "251662336", "behindDoc": "1", "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for axis, relative_from, offset in (("H", "column", from_left_margin), ("V", "page", from_page_top)):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", relative_from)
        position_offset = OxmlElement("wp:posOffset")
        position_offset.text = str(offset)
        position.append(position_offset)
        anchor.append(position)
    for child in list(inline):  # extent, docPr, cNvGraphicFramePr, graphic -- wrapping goes before docPr
        if child.tag == qn("wp:docPr"):
            anchor.append(OxmlElement("wp:wrapNone"))
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def _rule(paragraph, side: str, *, space_pt: int = 4):
    """A blue 1.5pt rule along one side of the paragraph, as under the title block / above the footer."""
    border = OxmlElement(f"w:{side}")
    for key, value in {"w:val": "single", "w:sz": "12", "w:space": str(space_pt), "w:color": _BM_BLUE}.items():
        border.set(qn(key), value)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    successors = (
        "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
        "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )
    paragraph._p.get_or_add_pPr().insert_element_before(borders, *successors)


def _add_field(paragraph, instruction: str):
    """Appends a live Word field (python-docx has no API for field codes)."""
    for kind in ("begin", "instr", "separate", "end"):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instruction
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def _add_page_of_total(paragraph):
    _add_field(paragraph, " PAGE ")
    paragraph.add_run("/")
    _add_field(paragraph, " NUMPAGES ")


def _footer(footer):
    line = footer.paragraphs[0]
    _rule(line, "top", space_pt=6)
    line.paragraph_format.tab_stops.add_tab_stop(_CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    _add_logo(line, _LOGO_MARK)
    line.add_run("   Barge Master\t")
    _add_page_of_total(line)


def _apply_header_footer(document, running_header: str, date_label: str):
    section = document.sections[0]
    section.different_first_page_header_footer = True
    for part in (section.first_page_header, section.first_page_footer, section.header, section.footer):
        part.is_linked_to_previous = False

    # Title page: the full logo floats top right, beside the title block (sized and placed as in
    # the Memo template); the following pages carry "Ref. - Subject" and the date instead.
    logo = _add_logo(section.first_page_header.paragraphs[0], _LOGO_FULL, width=1649730, height=901065)
    _float(logo, from_left_margin=4225262, from_page_top=810000)  # level with the title
    section.header.paragraphs[0].style = "BM Running Header"
    section.header.paragraphs[0].text = running_header
    section.header.add_paragraph(date_label, style="Header")
    _footer(section.first_page_footer)
    _footer(section.footer)


# ---- Content helpers ----

_SUBSCRIPTS = re.compile(r"\b([TH])_([ps])\b")  # T_p, H_s -> Tp, Hs with the p/s subscripted


def _add_text(paragraph, text: str, *, bold: bool = False, color: RGBColor | None = None):
    def run(part: str, subscript: bool = False):
        if not part:
            return
        r = paragraph.add_run(part)
        if bold:
            r.bold = True
        if color is not None:
            r.font.color.rgb = color
        if subscript:
            r.font.subscript = True

    position = 0
    for match in _SUBSCRIPTS.finditer(text):
        run(text[position : match.start()] + match.group(1))
        run(match.group(2), subscript=True)
        position = match.end()
    run(text[position:])


def _compose(document, *parts, style: str | None = None):
    """A paragraph of standard text with [prompts] for the author where it needs their input."""
    paragraph = document.add_paragraph(style=style)
    for part in parts:
        if isinstance(part, _Prompt):
            r = paragraph.add_run(f"[{part}]")
            r.italic = True
            r.font.color.rgb = _PLACEHOLDER_COLOR
        else:
            _add_text(paragraph, part)
    return paragraph


def _placeholder(document, text: str):
    return _compose(document, _Prompt(text))


def _shade(cell, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _fixed_widths(table, widths):
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width  # grid widths, which LibreOffice uses
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width  # per-cell widths, which Word uses


def _table_caption(document, text: str):
    caption = _compose(document, text, style="Caption")
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True


def _bm_table(document, header: list[str], rows: list[list[str]]):
    """BM's table look: solid blue header row with white bold text, light gray body, no grid."""
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    _fixed_widths(table, [Emu(_CONTENT_WIDTH // len(header))] * len(header))
    for row_index, values in enumerate([header, *rows]):
        for col, text in enumerate(values):
            cell = table.cell(row_index, col)
            cell.paragraphs[0].style = "BM Table Text"
            _add_text(cell.paragraphs[0], text, bold=row_index == 0, color=_WHITE if row_index == 0 else None)
            _shade(cell, _HEADER_FILL if row_index == 0 else _BODY_FILL)
    # The header row repeats if the table runs onto a next page; a short table isn't split at all.
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    if len(rows) <= 15:
        for row in table.rows[:-1]:
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True


def _title_block(document, rows: list[tuple[str, str]]):
    table = document.add_table(rows=len(rows), cols=2)
    _fixed_widths(table, (Cm(2.5), Cm(8.5)))  # clear of the logo on the right
    margins = OxmlElement("w:tblCellMar")
    for side in ("w:left", "w:right"):
        margin = OxmlElement(side)
        margin.set(qn("w:w"), "0")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table._tbl.tblPr.insert_element_before(margins, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    for (label, value), row in zip(rows, table.rows):
        for cell, text in zip(row.cells, (label, value)):
            cell.paragraphs[0].style = "BM Title Block"
            if text:
                cell.paragraphs[0].add_run(text)
    rule = document.add_paragraph()
    rule.paragraph_format.line_spacing = Pt(6)
    rule.paragraph_format.space_after = Pt(8)
    _rule(rule, "bottom")


def _assumption_lines(document, rows: list[tuple[str, str]]):
    for index, (label, value) in enumerate(rows):
        line = document.add_paragraph()
        line.paragraph_format.tab_stops.add_tab_stop(Cm(4.75))
        if index < len(rows) - 1:
            line.paragraph_format.space_after = Pt(0)
        _add_text(line, f"{label}:\t{value}")


def _picture(image: Path | bytes, max_width_cm: float, max_height_cm: float) -> tuple[io.BytesIO, Cm, Cm]:
    data = image if isinstance(image, bytes) else image.read_bytes()
    with Image.open(io.BytesIO(data)) as img:
        width_px, height_px = img.size
        dpi = float(img.info.get("dpi", (96, 96))[0] or 96)
        stream = io.BytesIO()
        if img.format in ("PNG", "JPEG", "GIF"):
            stream.write(data)
        else:  # python-docx can't embed e.g. WebP, so re-encode anything else as PNG
            img.convert("RGBA").save(stream, format="PNG")
    width_cm, height_cm = width_px / dpi * 2.54, height_px / dpi * 2.54
    scale = min(1.0, max_width_cm / width_cm, max_height_cm / height_cm)
    stream.seek(0)
    return stream, Cm(width_cm * scale), Cm(height_cm * scale)


def _frame(run):
    """A thin dark-blue frame around the run's picture, as BM's reports draw around their figures."""
    line = OxmlElement("a:ln")
    line.set("w", "9525")
    fill = OxmlElement("a:solidFill")
    color = OxmlElement("a:srgbClr")
    color.set("val", _FRAME_COLOR)
    fill.append(color)
    line.append(fill)
    run._r.find(".//" + qn("pic:spPr")).append(line)


def _figure(
    document, image: Path | bytes | None, caption: str, missing_prompt: str, *,
    max_width_cm: float = 15.0, framed: bool = True, centered: bool = True,
):
    """A picture (a file, or the PNG bytes of a plot made here) with its caption below."""
    alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    if isinstance(image, bytes) or (image is not None and image.exists()):
        stream, width, height = _picture(image, max_width_cm, 18.0)
        paragraph = document.add_paragraph()
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run()
        run.add_picture(stream, width=width, height=height)
        if framed:
            _frame(run)
    else:
        _placeholder(document, missing_prompt).paragraph_format.keep_with_next = True
    _compose(document, caption, style="Caption").alignment = alignment


# ---- Standard texts ----
#
# Each returns a list of paragraphs, each a list of text parts and [prompts] (see _compose).


def _write_paragraphs(document, paragraphs: list[list[str]]):
    for parts in paragraphs:
        _compose(document, *parts)


def _variant(machine_type: str) -> tuple[str | None, str | None]:
    """(operation, load) of the PyBM machine variant the study was run with: ("transfer", None),
    ("lifting", "highload" | "lowload"), or (None, None) for systems without variants."""
    if machine_type.endswith("_transfer"):
        return "transfer", None
    for load in _LIFT_CAPACITY_KG:
        if machine_type.endswith(f"_{load}"):
            return "lifting", load
    return None, None


def _side(machine: dict) -> str | None:
    return {90: "port side", 270: "starboard side"}.get(machine.get("slewing_angle_deg"))


def _join(parts: list[str]) -> str:
    return parts[0] if len(parts) == 1 else ", ".join(parts[:-1]) + " and " + parts[-1]


def _overview_rows(result: dict) -> list[tuple[str, dict]]:
    """Head seas and the +/-30 degree headings, plus the worst-case direction if it's another one."""
    rows = {row["heading_from_deg"]: row for row in result.get("table") or []}
    overview = [(label, rows[heading]) for heading, label in _SUMMARY_HEADINGS if heading in rows]
    worst = min(rows.values(), key=lambda row: row["hs_limit_min_m"], default=None)
    if worst is not None and worst["heading_from_deg"] not in {heading for heading, _ in _SUMMARY_HEADINGS}:
        overview.append((f"Worst-case direction ({worst['heading_from_deg']}°)", worst))
    return overview


def _hs_at_tp(row: dict) -> str:
    """'2.0 m Hs @ Tp = 10 s': a heading's lowest limiting Hs and the period it occurs at."""
    text = f"{row['hs_limit_min_m']:.1f} m H_s"
    if row.get("tp_at_min_s") is not None:
        text += f" @ T_p = {_num(row['tp_at_min_s'])} s"
    return text


def _executive_summary(result: dict | None, location: str | None) -> list[list[str]]:
    if result is None:
        return [[_Prompt(
            "Summarize the results once a result is approved: the limiting Hs for head seas and the ±30° headings, "
            "the most critical direction, and what this means for the operation."
        )]]
    rows = {row["heading_from_deg"]: row for row in result.get("table") or []}
    limits = [
        f"approximately {rows[heading]['hs_limit_min_m']:.1f} m H_s for {name}"
        for heading, name in ((180, "head seas"), (210, "a +30-degree heading"), (150, "a -30-degree heading"))
        if heading in rows
    ]
    text = f"The calculations show a minimum workability of {_join(limits)}." if limits else ""
    worst = min(rows.values(), key=lambda row: row["hs_limit_min_m"], default=None)
    if worst is not None and worst["heading_from_deg"] not in (150, 180, 210):
        text += (
            f" The most critical direction ({worst['heading_from_deg']}°) corresponds to a workability limit of "
            f"{worst['hs_limit_min_m']:.1f} m H_s."
        )
    if result.get("overall_workability_pct") is not None:
        text += (
            f" Using the metocean data of {location}, and considering the worst-case wave direction, the yearly average "
            f"workability is {_num(result['overall_workability_pct'])}%."
        )
    text += " Detailed results per incoming wave direction and peak wave period (T_p) can be found in Appendix A. "
    return [[text.lstrip(), _Prompt(_PROMPTS["executive_summary"])]]


def _introduction(system_type: str, full_name: str, vessel: str, operation: str | None) -> list[list[str]]:
    if system_type == "t700":
        return [
            [
                f"The objective of this document is to present the workability of the {full_name} for ",
                _Prompt("lifting operations"),
                f" on the {vessel}. The motions at the location of the BM-T700 system on the {vessel} are compared to "
                "the capabilities of a BM-T700 in terms of actuator stroke, velocity and platform heave velocity. From "
                "this, the maximum sea state for safe working is calculated. This analysis is based on the specific setup "
                "on the ship and environmental conditions for this project.",
            ],
            [
                "The BM-T700 compensates three degrees of freedom (DoF); one translation (heave) and two rotations (roll "
                "and pitch) by controlling the 3 vertical hydraulic cylinders such that the platform is kept in an earth "
                "fixed steady position. The remaining three DoFs, two translations (surge and sway) and one rotation "
                "(yaw), are constrained by means of constraining rods. These motions are not compensated."
            ],
        ]
    what = "the workability of the lifting operation" if operation == "lifting" else "the workability"
    return [[f"This document presents {what} that is to be expected with a {full_name} installed on the {vessel}."]]


def _results_basis(system_type: str, operation: str | None, machine: dict) -> list[list[str]]:
    if system_type != "mcg30":
        return [[_Prompt(_PROMPTS["introduction_narrative"])]]
    luffing = machine.get("luffing_angle_deg")
    angle = f"a luffing angle of {_num(luffing)} degrees" if luffing is not None else _Prompt("luffing angle")
    if operation == "lifting":
        return [[
            "Results are presented for the ", _Prompt("standard"), " pedestal, with lifting workability calculated for ",
            angle, f" and a lifting radius of {_num(machine.get('lifting_radius_m'))} m.",
        ]]
    return [[
        "Results are presented for the ", _Prompt("standard"), " pedestal, with personnel transfer workability calculated for ",
        angle, ", to accommodate a landing height of ", _Prompt("+xx.x m MSL"), ".",
    ]]


def _general_arrangement(system_type: str, label: str, vessel: str, operation: str | None, figure: int) -> list[list[str]]:
    if system_type == "t700":
        return [[
            f"The BM-T700 system will be installed on the {vessel}. The general arrangement (GA) of this vessel with the "
            f"BM-T700 is shown in Figure {figure}."
        ]]
    paragraphs = [[f"The location of the {label} onboard the {vessel} is shown in Figure {figure}."]]
    if operation == "lifting":
        paragraphs.append([_Prompt(
            "Describe the hoisting capabilities relative to the MCG COG: the minimum and maximum hoisting range, and the "
            "resulting distance between the vessel and the drop location."
        )])
    return paragraphs


# The BM-T700 technical data its workability reports state (stroke and velocities as in PyBM's T700 class).
_T700_SPECIFICATION = (
    "Maximum cylinder stroke: 2.54 m",
    "Operational cylinder stroke: 2.20 m",
    "Operational cylinder speed: 0.74 m/s",
    "Maximum residual platform velocity: 0.44 m/s",
)
_T700_MOTION_MONITOR = (
    "The BM-T700 does not need integration with the vessel’s Deck Motion Monitor, it has stand-alone motion sensors and "
    "is actuated without input from the vessel’s Deck Motion Monitor. However, options may be explored to use the "
    "vessel’s Deck Motion Monitor as indicative tool during the operation."
)


def _simulation_parameters(
    study: dict, system_type: str, label: str, operation: str | None, load: str | None, machine: dict, table: int
) -> list[list[str]]:
    draft = study.get("vessel_draft_m")
    draught = f" The draught of the vessel in the considered loading condition is {_num(draft)} m." if draft else ""
    if system_type != "mcg30":
        return [[
            f"The input parameters are given in Table {table}. The convention for wave directions used in the table "
            f"corresponds with Figure 1.{draught}"
        ]]
    side = _side(machine)
    parts = [
        f"The workability estimation is performed for one scenario. The scenario assumes a {label} on the ",
        _Prompt("standard"),
        f" pedestal{f' with {side} operations' if side else ''}. All parameters, including the vessel and wave spectrum, "
        f"remain constant.{draught}",
    ]
    if operation == "lifting":
        parts.append(
            " From the lifting arrangement, the luffing angle and lifting radius are determined. Furthermore, this study "
            f"is based on an engineered lift of up to {_LIFT_CAPACITY_KG[load]} kg. This information is also shown in "
            f"Table {table}."
        )
    return [parts]


def _results_text(label: str, machine: dict) -> list[list[str]]:
    side = _side(machine)
    return [[
        f"The workability estimations of the {label} are provided below{f' for {side} operations' if side else ''}. "
        "Results for bow waves and ±30-degree headings are shown below. Results for other wave directions can be found "
        "in Appendix A. ",
        _Prompt(_PROMPTS["estimation_narrative"]),
    ]]


def _criterion(system_type: str, short: str, n_sigma: float | None) -> list[list[str]]:
    if n_sigma is not None:
        basis = (
            f"For this workability study, the limit is defined such that {_num(n_sigma)} standard deviations of the "
            "stochastic movements should stay below the system limits."
        )
    else:
        basis = (
            "For this workability study, a “Most Probable Extreme” is calculated, for an expected operation time of "
            "3 hours. This means that the probability of exceeding any system limit is such that we expect that it will "
            "not occur in a period of 3 hours."
        )
    limit = (
        "The workability limit is defined as the significant wave height that results in a certain probability of "
        f"exceeding a system limit, such as cylinder stroke. {basis} Due to the probabilistic nature of waves and "
        "movements, there is still a possibility that a wave occurs that will exceed the limit of the machine."
    )
    if system_type == "t700":
        return [
            [limit + " These extreme waves usually exceed the capability of the system by only a few percent causing "
             "minimal residual motion, without shocks."],
            ["If the BM-T700 encounters such an extreme wave, a sophisticated “compensation limiter” software function "
             "automatically becomes active that temporarily reduces the heave compensation. This protects the physical "
             "limits of the system, while ensuring a smooth and minimized residual platform heave motion for a safe and "
             "predictable operation. Roll and pitch motions are always fully compensated. After the high vessel motion, "
             "full motion compensation continues automatically."],
        ]
    return [
        ["Due to the stochastic nature of waves and resulting vessel motions, it is not feasible to design a motion "
         "compensation system that allows full compensation of all vessel motions occurring during all operational "
         "conditions. For example, increasing the cylinder stroke will reduce the probability that the vessel motion "
         "exceeds the stroke for a given sea-state, but will never reduce the probability to zero."],
        [limit],
        [f"If the {short} detects such an extreme wave, a software function detects that an exceedance of limits may "
         "occur. This function will alarm the operator and crew so that they can respond accordingly."],
    ]


# ---- Calculation data ----


def _num(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _input_parameter_rows(study: dict, version: dict, short: str, load: str | None, refs: dict) -> list[list[str]]:
    env = version["params"].get("environment", {})
    m = version["params"].get("machine", {})
    vessel, draft = study["vessel"] or "—", study.get("vessel_draft_m")
    rows = [
        ["Vessel", f"{vessel} (draught = {_num(draft)} m) [{refs['rao']}]" if draft else f"{vessel} [{refs['rao']}]"],
        ["Loading condition", study["loadcase"] or "—"],
        [
            "Wave Spectrum",
            f"Jonswap, Directional spreading n={_num(env.get('spreading'))}, Gamma according to DNV regulations [{refs['dnv']}]",
        ],
        ["T_p", f"{_num(env.get('tp_min_s'))}-{_num(env.get('tp_max_s'))} sec"],
        ["Incoming wave headings", f"0-360 degrees, {_num(env.get('wave_dir_step_deg'))}-degree step"],
        [
            f"{short} Location (aft, centreline, keel)",
            f"[{_num(m.get('position_x_m'))} {_num(m.get('position_y_m'))} {_num(m.get('position_z_m'))}] m",
        ],
    ]
    # Only the machine fields this version's system type actually used (see pybm_service.run_calculation).
    if m.get("luffing_angle_deg") is not None:
        rows.append([f"{short} Luffing Angle", f"{_num(m['luffing_angle_deg'])} deg"])
    if m.get("lifting_radius_m") is not None:
        rows.append([f"{short} Lifting radius", f"{_num(m['lifting_radius_m'])} m"])
    if m.get("slewing_angle_deg") is not None:
        side = _side(m)
        rows.append([f"{short} Slewing angle", f"{_num(m['slewing_angle_deg'])} degrees{f' (i.e., toward {side})' if side else ''}"])
    if load:
        rows.append(["Engineered lift maximum weight", f"{_LIFT_CAPACITY_KG[load]} kg"])
    if m.get("dp_accuracy_m") is not None:
        rows.append(["DP2 Accuracy (+/-)", f"{_num(m['dp_accuracy_m'])} m"])
    if m.get("orientation_deg") is not None:
        rows.append([f"{short} Orientation", f"{_num(m['orientation_deg'])} degrees"])
    if m.get("p_passive_pa") is not None:
        rows.append(["Passive pressure", f"{_num(round(m['p_passive_pa'] / 1e5, 1))} bar"])
    if m.get("comp_point_x_m") is not None:
        rows.append(
            [
                "Compensation point (w.r.t. central bearing)",
                f"[{_num(m['comp_point_x_m'])} {_num(m.get('comp_point_y_m'))} {_num(m.get('comp_point_z_m'))}] m",
            ]
        )
    if "metocean" in refs:
        rows.append(["Metocean data", f"{version['result']['location_label']} [{refs['metocean']}]"])
    rows.append(["Workability criterion", "As given in Appendix B"])
    return rows


# Head seas in BM blue, the +/-30 degree headings in navy and gray, the worst-case direction in red.
_PLOT_LINES = ((180, "Head seas (180°)", "#00B0F0"), (210, "+30° heading (210°)", "#002060"), (150, "-30° heading (150°)", "#7F7F7F"))


def _hs_tp_plot(result: dict) -> bytes | None:
    """Limiting Hs against Tp for bow waves, the +/-30 degree headings and the worst-case direction --
    the results figure of BM's reports. Needs the per-Tp curves, which results stored before they
    were kept don't have."""
    tps = result.get("tps_s")
    rows = {row["heading_from_deg"]: row for row in result.get("table") or [] if row.get("hs_limit_by_tp_m")}
    if not tps or not rows:
        return None
    from matplotlib.figure import Figure  # no pyplot: no GUI backend involved

    lines = [(heading, label, color) for heading, label, color in _PLOT_LINES if heading in rows]
    worst = min(rows.values(), key=lambda row: row["hs_limit_min_m"])
    if worst["heading_from_deg"] not in {heading for heading, _, _ in lines}:
        lines.append((worst["heading_from_deg"], f"Worst-case direction ({worst['heading_from_deg']}°)", "#E03C31"))
    figure = Figure(figsize=(6.3, 3.6), dpi=200)
    axes = figure.add_subplot()
    for heading, label, color in lines:
        axes.plot(tps, rows[heading]["hs_limit_by_tp_m"], marker="o", markersize=3, linewidth=1.6, color=color, label=label)
    axes.set_xlabel("$T_p$ [s]")
    axes.set_ylabel("Limiting $H_s$ [m]")
    axes.set_ylim(bottom=0)
    axes.grid(alpha=0.3)
    axes.legend(fontsize=8)
    figure.tight_layout()
    out = io.BytesIO()
    figure.savefig(out, format="png")
    return out.getvalue()


def today_label() -> str:
    today = date.today()
    return f"{today.day}-{today.month}-{today.year}"


def _header_fields(study: dict, report: dict, system_label: str, author: str) -> dict:
    defaults = {
        "subject": f"Workability report of the {system_label} onboard the {study['vessel'] or '[vessel]'}",
        "project": study["project"] or "",
        "ref": "",
        "date": today_label(),
    }
    saved = report.get("header") or {}
    fields = {key: str(saved.get(key) or default) for key, default in defaults.items()}
    # Author is whoever generates the document; Check and Approval are the reviewer, set on submission.
    return {**fields, "author": author, "check": "", "approval": ""}


class _Counter:
    """Figure and table numbers, in the order they appear."""

    def __init__(self):
        self.figure = self.table = 0

    def next_figure(self) -> int:
        self.figure += 1
        return self.figure

    def next_table(self) -> int:
        self.table += 1
        return self.table


def build_report_docx(study: dict, author: str) -> bytes:
    """Fills BM's report layout with everything the app knows plus BM's standard texts;
    prompts in [brackets] mark what's left to write."""
    report = study.get("report") or {}
    system_type = study["system_type"]
    label = SYSTEM_LABELS.get(system_type, system_type or "system")
    full_name = _SYSTEM_FULL_NAMES.get(system_type, label)
    short = _SYSTEM_SHORT_NAMES.get(system_type, label)
    vessel = study["vessel"] or "[vessel]"
    header = _header_fields(study, report, label, author)
    approved = next((v for v in study["versions"] if v["approved"]), None)
    result = approved["result"] if approved else None
    machine = approved["params"].get("machine", {}) if approved else {}
    operation, load = _variant(machine.get("machine_type", ""))
    location = result.get("location_label") if result else None
    images = {f["id"]: f for f in study["files"] if f["kind"] == "image"}
    numbers = _Counter()

    def image_path(file_id: str | None) -> Path | None:
        f = images.get(file_id) if file_id else None
        return db.STUDY_FILES_DIR / f["filename"] if f else None

    document = Document()
    _apply_house_style(document)
    running_header = f"{header['ref']} - {header['subject']}" if header["ref"] else header["subject"]
    _apply_header_footer(document, running_header, header["date"])
    document.core_properties.title = header["subject"]
    document.core_properties.subject = header["subject"]
    document.core_properties.category = header["ref"]
    document.core_properties.author = header["author"] or study.get("owner_name") or ""

    title = "Workability estimation lifting operation" if operation == "lifting" else "Workability report"
    document.add_paragraph(title, style="BM_H0")
    _title_block(
        document,
        [("Subject", header["subject"]), ("Project", header["project"]), ("Ref.", header["ref"]), ("Date", header["date"]),
         ("Author", header["author"]), ("Check", header["check"]), ("Approval", header["approval"])],
    )

    # 1 -- with the overview of results that opens every BM report
    document.add_paragraph("1\tExecutive summary", style="Heading 1")
    _write_paragraphs(document, _executive_summary(result, location))
    overview = _overview_rows(result) if result else []
    if overview:
        _table_caption(document, f"Table {numbers.next_table()}: Overview of results")
        _bm_table(document, ["Wave heading", "Workability limit"], [[name, _hs_at_tp(row)] for name, row in overview])

    # 2
    document.add_paragraph("2\tIntroduction", style="Heading 1").paragraph_format.page_break_before = True
    _write_paragraphs(document, _introduction(system_type, full_name, vessel, operation))
    _compose(document, "For this workability estimation, BM has taken the following assumptions:")
    _assumption_lines(
        document,
        [("Vessel", vessel), ("Loading condition", study["loadcase"] or "—")]
        + ([("Area/environmental data", location)] if location else []),
    )
    _write_paragraphs(document, _results_basis(system_type, operation, machine))
    _compose(document, "The system of axis and naming conventions used are presented in Figure 1.")
    _figure(
        document,
        image_path(report.get("axis_convention_file_id")) or _AXIS_CONVENTION,
        f"Figure {numbers.next_figure()}: System of axis, sign and naming conventions. Z-coordinate of the vessel is zero at the keel.",
        "Insert the axis/sign convention diagram here",
        max_width_cm=10.0,
    )

    # 3
    document.add_paragraph("3\tInput parameters", style="Heading 1")
    document.add_paragraph("General arrangement", style="Heading 2")
    figure = numbers.next_figure()
    _write_paragraphs(document, _general_arrangement(system_type, label, vessel, operation, figure))
    _figure(
        document,
        image_path(report.get("concept_drawing_file_id")),
        f"Figure {figure}: General arrangement of the {label} onboard the {vessel}.",
        "Insert the general arrangement / deck layout showing the location of the system here",
    )
    extra = [images[file_id] for file_id in report.get("additional_figure_ids") or [] if file_id in images]
    if extra:
        document.add_paragraph("Additional information", style="Heading 2")
        for f in extra:
            _figure(document, db.STUDY_FILES_DIR / f["filename"], f"Figure {numbers.next_figure()}: {f['caption'] or f['original_filename']}", "")
    if system_type == "t700":
        document.add_paragraph("BM-T700 system specification", style="Heading 2")
        _compose(document, "The operational envelope analysis is based on the following Barge Master technical data:")
        for index, item in enumerate(_T700_SPECIFICATION):
            bullet = _compose(document, item, style="List Bullet")
            if index < len(_T700_SPECIFICATION) - 1:
                bullet.paragraph_format.space_after = Pt(0)
        _compose(document, _T700_MOTION_MONITOR)
    document.add_paragraph("Simulation parameters", style="Heading 2")
    refs = {"dnv": 1, "rao": 2, **({"metocean": 3} if location else {})}
    if approved:
        table = numbers.next_table()
        _write_paragraphs(document, _simulation_parameters(study, system_type, label, operation, load, machine, table))
        _table_caption(document, f"Table {table}: Input parameters")
        _bm_table(document, ["Parameter", "Value"], _input_parameter_rows(study, approved, short, load, refs))
    else:
        _placeholder(
            document,
            "No approved result yet — approve one in Build study, then regenerate this report to fill in the input "
            "parameters and results.",
        )

    # 4
    document.add_paragraph("4\tWorkability results", style="Heading 1").paragraph_format.page_break_before = True
    plot = db.PLOTS_DIR / Path(approved["plot_url"]).name if approved and approved.get("plot_url") else None
    hs_tp = _hs_tp_plot(result) if result else None
    if approved:
        _write_paragraphs(document, _results_text(label, machine))
        what = {
            "transfer": " for personnel transfer operation",
            "lifting": f" for {'high' if load == 'highload' else 'low'} load lifting operation",
        }.get(operation, "")
        _figure(
            document,
            hs_tp or plot,
            f"Figure {numbers.next_figure()}: {short} Workability{what}.",
            "No result plot was stored for the approved result",
            max_width_cm=15.0 if hs_tp else 12.0,
        )
        if result.get("overall_workability_pct") is not None:
            document.add_paragraph("Workability percentages", style="Heading 2")
            _compose(
                document,
                f"Using metocean data of {location} [{refs['metocean']}], the limiting H_s is translated into a percentage "
                "of workable sea states. For this, the worst-case wave direction at each wave period was considered. The "
                f"yearly average workability for this worst-case scenario is {_num(result['overall_workability_pct'])}%.",
            )
            monthly = result.get("monthly_workability_pct")
            if monthly:
                _table_caption(document, f"Table {numbers.next_table()}: Monthly workability")
                _bm_table(document, ["Month", "Workability"], [[month, f"{pct:.1f}%"] for month, pct in monthly.items()])
    else:
        _placeholder(document, "The workability results follow once a result is approved.")

    # 5
    document.add_paragraph("5\tReferences", style="Heading 1")
    references = [
        [_DNV_REFERENCE],
        [f"Motion RAOs of the {vessel}, loading condition {study['loadcase'] or '—'}: ", _Prompt("RAO document, received on date")],
    ] + ([[f"Metocean data: {location}"]] if location else [])
    for number, parts in enumerate(references, start=1):
        entry = _compose(document, f"[{number}]\t", *parts)
        entry.paragraph_format.tab_stops.add_tab_stop(Cm(1))
        entry.paragraph_format.left_indent = Cm(1)
        entry.paragraph_format.first_line_indent = Cm(-1)
        entry.paragraph_format.space_after = Pt(4)

    # Appendix A -- the full envelope: the rose plot and the limit per heading
    if approved:
        appendix = document.add_paragraph("Appendix A: Results", style="BM_Appendix")
        appendix.paragraph_format.page_break_before = True
        section = 1
        if hs_tp and plot is not None:
            _compose(document, f"A.{section}\tResults in roseplot format", style="Heading 2")
            _figure(document, plot, f"Figure {numbers.next_figure()}: {short} Workability per wave period in roseplot format.", "", max_width_cm=12.0)
            section += 1
        _compose(document, f"A.{section}\tLimiting H_s per wave heading", style="Heading 2")
        _table_caption(document, f"Table {numbers.next_table()}: Limiting H_s per incoming wave heading")
        _bm_table(
            document,
            ["Wave heading", "Minimum H_s [m]", "At T_p [s]", "Maximum H_s [m]"],
            [
                [f"{row['heading_from_deg']}°", f"{row['hs_limit_min_m']:.2f}", _num(row.get("tp_at_min_s")), f"{row['hs_limit_max_m']:.2f}"]
                for row in result.get("table") or []
            ],
        )

    # Appendix B
    appendix = document.add_paragraph("Appendix B: Workability criterion", style="BM_Appendix")
    appendix.paragraph_format.page_break_before = True
    n_sigma = result.get("n_sigma_used") if result else None
    _write_paragraphs(document, _criterion(system_type, short, n_sigma))

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---- Title block updates on an existing (possibly hand-edited) document ----


def _title_block_table(document):
    """The Subject/Project/Ref./... table -- found by its labels, since the author may have
    added content above it."""
    for table in document.tables:
        labels = {row.cells[0].text.strip() for row in table.rows if len(row.cells) >= 2}
        if {"Subject", "Check"} <= labels:
            return table
    return None


def _set_paragraph_text(paragraph, text: str):
    """Replaces the paragraph's text, keeping the first run's formatting."""
    runs = list(paragraph.runs)
    run_format = deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    for run in runs:
        run._r.getparent().remove(run._r)
    new = paragraph.add_run(text)
    if run_format is not None:
        new._r.insert(0, run_format)


def title_block_values(docx: bytes) -> dict[str, str]:
    table = _title_block_table(Document(io.BytesIO(docx)))
    if table is None:
        return {}
    return {row.cells[0].text.strip(): row.cells[1].text.strip() for row in table.rows if len(row.cells) >= 2}


def _update_running_header(document, label: str, old: str, new: str, subject: str):
    """The running header repeats the Ref. and Date: a "Ref. - Subject" line (just the subject
    while there's no Ref. yet) and a line with the date."""
    for section in document.sections:
        if section.header.is_linked_to_previous:
            continue  # touching it would add a header the document doesn't have
        for paragraph in section.header.paragraphs:
            text = paragraph.text.strip()
            if label == "Date":
                if old and text == old:
                    _set_paragraph_text(paragraph, new)
            elif old and text.startswith(old):
                _set_paragraph_text(paragraph, new + text[len(old) :])
            elif not old and new and subject and text == subject:
                _set_paragraph_text(paragraph, f"{new} - {subject}")


def update_title_block(docx: bytes, fields: dict[str, str]) -> bytes:
    """Sets title-block values by label (e.g. {"Check": "J. Doe"}); labels the author removed
    are skipped. A new Ref. or Date also replaces the old one in the running page header."""
    document = Document(io.BytesIO(docx))
    table = _title_block_table(document)
    if table is None:
        return docx
    values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in table.rows if len(row.cells) >= 2}
    for row in table.rows:
        if len(row.cells) < 2 or row.cells[0].text.strip() not in fields:
            continue
        label, value_cell = row.cells[0].text.strip(), row.cells[1]
        old = value_cell.text.strip()
        _set_paragraph_text(value_cell.paragraphs[0], fields[label])
        for extra in value_cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        if label in ("Ref.", "Date") and fields[label] != old:
            _update_running_header(document, label, old, fields[label], values.get("Subject", ""))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def bump_ref_revision(ref: str, revision: int) -> str:
    """'BM-1234-EN-0001-Rev0' -> '...-Rev1'. A ref without a trailing Rev number is left as is."""
    return re.sub(r"(rev\s*)\d+$", lambda m: f"{m.group(1)}{revision}", ref.strip(), flags=re.IGNORECASE)
