"""Review comments as real Word comments -- for rewriting a report in Word after its review
(ReviewComments.tsx's "Download for Word, with comments") -- and taking them out again
(strip_comments) when that document comes back into the app: the app's own comment threads
stay the record, so Word comments never end up in a report or doubled on the next download.

The app's threads (db.py's review_comments) point at their text by paragraph id plus the quoted
words (docx_anchors.py). Each thread becomes a Word comment on that text, its replies become
Word replies, and a resolved thread is marked done. python-docx writes the comments themselves;
replies and the done state live in commentsExtended.xml, which python-docx doesn't know, so
that part is added here.

Where a thread lands, best first:
  1. the quoted words, in the paragraph with its id (runs are split at the quote's edges);
  2. that whole paragraph, when the words have since changed;
  3. the quoted words anywhere, when the paragraph itself is gone (a later revision);
  4. the document's first paragraph, with the quote written into the comment -- also where
     general comments (no quote) go.
"""

import copy
import io
import re
import secrets
import zipfile
from datetime import datetime, timezone

from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

_W14_PARA_ID = "{http://schemas.microsoft.com/office/word/2010/wordml}paraId"
_W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
_EXTENDED_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"
_EXTENDED_REL = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
# Run content that can be split as plain text without losing anything that matters.
_PLAIN_RUN_CHILDREN = {qn("w:rPr"), qn("w:t"), qn("w:lastRenderedPageBreak")}
# Every part Word keeps comments in: the comments themselves, threads and done states, the ids
# modern Word adds, and its comment extensions.
_COMMENT_RELS = {
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
    _EXTENDED_REL,
    "http://schemas.microsoft.com/office/2016/09/relationships/commentsIds",
    "http://schemas.microsoft.com/office/2018/08/relationships/commentsExtensible",
}
_COMMENT_MARKERS = {qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:commentReference")}
_DOCUMENT_RELS = "word/_rels/document.xml.rels"
_PKG_TYPES = "{http://schemas.openxmlformats.org/package/2006/content-types}"


def strip_comments(docx: bytes) -> tuple[bytes, int]:
    """`docx` without any Word comments -- their markers in the text (body, headers, footers,
    notes), their parts and those parts' relationships and content types -- and how many
    comments there were. Everything else is copied byte for byte; a document without comments
    comes back unchanged."""
    with zipfile.ZipFile(io.BytesIO(docx)) as src:
        names = set(src.namelist())
        if _DOCUMENT_RELS not in names:
            return docx, 0
        rels = etree.fromstring(src.read(_DOCUMENT_RELS))
        comment_rels = [r for r in rels if r.get("Type") in _COMMENT_RELS]
        story_parts = [
            n for n in names
            if n.startswith("word/") and n.endswith(".xml") and "/" not in n[5:]
            and re.match(r"(document|header\d*|footer\d*|footnotes|endnotes)\.xml$", n[5:])
        ]
        marked = {n for n in story_parts if b"commentR" in src.read(n)}
        if not comment_rels and not marked:
            return docx, 0

        dropped = {"word/" + r.get("Target").lstrip("/").removeprefix("word/") for r in comment_rels}
        count = 0
        if "word/comments.xml" in dropped and "word/comments.xml" in names:
            count = len(etree.fromstring(src.read("word/comments.xml")).findall(qn("w:comment")))
        for r in comment_rels:
            rels.remove(r)

        replaced: dict[str, bytes] = {_DOCUMENT_RELS: etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)}
        for name in marked:
            root = etree.fromstring(src.read(name))
            for marker in [el for el in root.iter() if el.tag in _COMMENT_MARKERS]:
                run = marker.getparent()
                run.remove(marker)
                # The reference sat in a run of its own: drop the run once it's empty.
                if run.tag == qn("w:r") and all(child.tag == qn("w:rPr") for child in run):
                    run.getparent().remove(run)
            replaced[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        if "[Content_Types].xml" in names:
            types = etree.fromstring(src.read("[Content_Types].xml"))
            for override in types.findall(f"{_PKG_TYPES}Override"):
                if override.get("PartName", "").lstrip("/") in dropped:
                    types.remove(override)
            replaced["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)

        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                if item.filename in dropped:
                    continue
                dst.writestr(item, replaced.get(item.filename, src.read(item.filename)))
    return out.getvalue(), count


def add_review_comments(docx: bytes, comments: list[dict]) -> bytes:
    """`docx` with the threads in `comments` (db.list_comments rows: top-level comments plus
    replies) added as Word comments."""
    document = Document(io.BytesIO(docx))
    paragraphs = [Paragraph(p, document._body) for p in document.element.body.iter(qn("w:p"))]
    replies: dict[str, list[dict]] = {}
    for c in comments:
        if c["parent_id"]:
            replies.setdefault(c["parent_id"], []).append(c)

    extended: list[tuple[str, str | None, bool]] = []  # (paraId, parent paraId, done)
    taken_ids = {p._p.get(_W14_PARA_ID) for p in paragraphs} - {None}
    for root in (c for c in comments if not c["parent_id"]):
        runs, note = _anchor(paragraphs, root)
        done = bool(root["resolved_at"])
        root_comment = document.add_comment(runs, text=note or root["body"].strip(), **_who(root))
        if note:
            root_comment.add_paragraph(root["body"].strip())
        _set_date(root_comment, root)
        root_para = _stamp_last_paragraph(root_comment, taken_ids)
        extended.append((root_para, None, done))
        previous_id = root_comment.comment_id
        for reply in replies.get(root["id"], []):
            reply_comment = document.comments.add_comment(text=reply["body"].strip(), **_who(reply))
            _set_date(reply_comment, reply)
            _mark_reply_range(document, previous_id, reply_comment.comment_id)
            previous_id = reply_comment.comment_id
            extended.append((_stamp_last_paragraph(reply_comment, taken_ids), root_para, done))

    if extended:
        _write_extended(document, extended)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---- Where a thread goes ----


def _phrase(quote: str) -> str:
    # The quote's first line: a selection can span paragraphs, its first line sits in one.
    return next((line.strip() for line in (quote or "").splitlines() if line.strip()), "")


def _anchor(paragraphs: list[Paragraph], comment: dict) -> tuple[list[Run], str | None]:
    """The runs a thread goes on, plus a note for the comment text when that isn't the quote."""
    phrase = _phrase(comment["quote"])
    quoted = f"“{comment['quote'].strip()}”"
    para_id = (comment["para_id"] or "").upper()
    home = next((p for p in paragraphs if (p._p.get(_W14_PARA_ID) or "").upper() == para_id), None) if para_id else None
    if home is not None:
        if phrase:
            runs = _quote_runs(home, phrase)
            if runs:
                return runs, None
        return _paragraph_runs(home), f"On {quoted}:" if phrase else None
    if phrase:
        for p in paragraphs:
            runs = _quote_runs(p, phrase)
            if runs:
                return runs, None
    first = next((p for p in paragraphs if p.text.strip()), paragraphs[0])
    if phrase:
        return _paragraph_runs(first), f"On {quoted} (no longer in the document):"
    return _paragraph_runs(first), "General comment:"


def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
    return paragraph.runs or [paragraph.add_run("")]


def _quote_runs(paragraph: Paragraph, phrase: str) -> list[Run] | None:
    """The runs covering exactly `phrase` in this paragraph (split to fit), or None if it isn't
    there. A run that can't be split safely (a tab, a field, a picture) is taken whole."""
    start = "".join(r.text for r in paragraph.runs).find(phrase)
    if start < 0:
        return None
    end = start + len(phrase)
    _split_at(paragraph, end)
    _split_at(paragraph, start)
    picked, pos = [], 0
    for run in paragraph.runs:
        length = len(run.text)
        if pos < end and pos + length > start:
            picked.append(run)
        pos += length
    return picked or None


def _split_at(paragraph: Paragraph, offset: int) -> None:
    pos = 0
    for run in paragraph.runs:
        length = len(run.text)
        if pos < offset < pos + length:
            if all(child.tag in _PLAIN_RUN_CHILDREN for child in run._r):
                text, cut = run.text, offset - pos
                twin = copy.deepcopy(run._r)
                run._r.addnext(twin)
                run.text = text[:cut]
                Run(twin, run._parent).text = text[cut:]
            return
        pos += length


# ---- Writing the comments ----


def _initials(name: str) -> str:
    return "".join(part[0] for part in re.split(r"[\s._-]+", name) if part)[:3].upper()


def _who(c: dict) -> dict:
    return {"author": c["author_name"], "initials": _initials(c["author_name"])}


def _set_date(comment, c: dict) -> None:
    try:
        comment._comment_elm.date = datetime.fromisoformat(c["created_at"]).astimezone(timezone.utc).replace(tzinfo=None)
    except (TypeError, ValueError):
        pass  # keeps python-docx's "now"


def _mark_reply_range(document, previous_id: int, reply_id: int) -> None:
    """A reply's range markers on the same text as its thread, each one just after the previous
    comment's -- Word only links a reply that comes after its parent in the document, and
    python-docx's own marking would nest it before."""
    body = document.element.body
    previous_start = body.xpath(f'.//w:commentRangeStart[@w:id="{previous_id}"]')[0]
    previous_reference = body.xpath(f'.//w:r[w:commentReference[@w:id="{previous_id}"]]')[0]
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(reply_id))
    previous_start.addnext(start)
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(reply_id))
    reference_run = copy.deepcopy(previous_reference)
    reference_run.xpath("./w:commentReference")[0].set(qn("w:id"), str(reply_id))
    previous_reference.addnext(end)
    end.addnext(reference_run)


def _stamp_last_paragraph(comment, taken: set[str]) -> str:
    """The w14:paraId of the comment's last paragraph -- how commentsExtended refers to it."""
    paragraph = comment._comment_elm.xpath("./w:p")[-1]
    while True:
        para_id = f"{secrets.randbelow(0x7FFFFFFF) + 1:08X}"
        if para_id not in taken:
            taken.add(para_id)
            paragraph.set(_W14_PARA_ID, para_id)
            return para_id


def _write_extended(document, entries: list[tuple[str, str | None, bool]]) -> None:
    """Replies (paraIdParent) and resolved threads (done) in word/commentsExtended.xml, added to
    the document's own part when it already has one (a file that's been through Word)."""
    existing = next((rel.target_part for rel in document.part.rels.values() if rel.reltype == _EXTENDED_REL), None)
    if existing is not None:
        root = etree.fromstring(existing.blob)
    else:
        root = etree.Element(f"{{{_W15}}}commentsEx", nsmap={"w15": _W15})
    for para_id, parent, done in entries:
        entry = etree.SubElement(root, f"{{{_W15}}}commentEx")
        entry.set(f"{{{_W15}}}paraId", para_id)
        if parent:
            entry.set(f"{{{_W15}}}paraIdParent", parent)
        entry.set(f"{{{_W15}}}done", "1" if done else "0")
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if existing is not None:
        existing._blob = blob
    else:
        part = Part(PackURI("/word/commentsExtended.xml"), _EXTENDED_TYPE, blob, document.part.package)
        document.part.relate_to(part, _EXTENDED_REL)
